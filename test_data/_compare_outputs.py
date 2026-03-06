"""Přímé porovnání dvou anon_output adresářů na úrovni anonymizovaných DOCX a map."""
import json
from pathlib import Path
from docx import Document
import re

DIR_A = Path(r'c:\Nixminds\skryi-clean\test_data\anon_output')
DIR_B = Path(r'c:\Nixminds\skryi_slozky\test_data\anon_output')
LABEL_A = 'skryi-clean'
LABEL_B = 'skryi_slozky'

IGNORE_PII = {
    'praha', 'brno', 'ostrava', 'plzeň', 'olomouc', 'liberec', 'pardubice',
    'hcp_admin', 'admin', 'bc.', 'mgr.', 'ing.', 'judr.', 'mudr.',
}


def analyze_output(anon_dir):
    results = {}
    for map_json in sorted(anon_dir.glob('*_map.json')):
        base = map_json.stem.replace('_map', '')
        if '_anon' in base or '_deanon' in base:
            continue
        anon_docx = anon_dir / f'{base}_anon.docx'
        map_txt = anon_dir / f'{base}_map.txt'
        if not anon_docx.exists():
            continue

        r = {'pii_leaks': [], 'phantom_tags': 0, 'txt_json_diff': 0, 'entity_count': 0, 'person_count': 0}

        try:
            doc = Document(anon_docx)
        except Exception as e:
            r['error'] = str(e)
            results[base] = r
            continue

        anon_text = ''
        for p in doc.paragraphs:
            anon_text += p.text + '\n'
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        anon_text += p.text + '\n'

        with open(map_json, 'r', encoding='utf-8') as f:
            map_data = json.load(f)

        labels_in_map = set()
        originals = {}
        for e in map_data.get('entities', []):
            orig = e.get('original', '')
            lbl = e.get('label', '')
            typ = e.get('type', '')
            if lbl:
                labels_in_map.add(lbl)
            if orig and len(orig) >= 4 and not orig.startswith('***REDACTED'):
                originals[orig] = typ

        r['entity_count'] = len(map_data.get('entities', []))
        r['person_count'] = sum(1 for e in map_data.get('entities', [])
                                if e.get('type') == 'PERSON' and e.get('label', '').endswith(']]'))

        # PII leaky - hodnoty z mapy nalezené v anonymu
        for orig, typ in originals.items():
            if orig.lower() in IGNORE_PII:
                continue
            if orig in anon_text:
                r['pii_leaks'].append((typ, orig[:60]))

        # Phantom tagy
        tags_in_text = set(re.findall(r'\[\[[A-Z_0-9]+\d*\]\]', anon_text))
        phantoms = tags_in_text - labels_in_map
        r['phantom_tags'] = len(phantoms)

        # TXT vs JSON
        if map_txt.exists():
            txt_labels = set()
            with open(map_txt, 'r', encoding='utf-8') as f:
                for line in f:
                    if line.strip().startswith('[['):
                        lbl = line.split(':')[0].strip()
                        txt_labels.add(lbl)
            r['txt_json_diff'] = len(txt_labels - labels_in_map) + len(labels_in_map - txt_labels)

        results[base] = r
    return results


print('Analyzuji oba adresare...\n')
res_a = analyze_output(DIR_A)
res_b = analyze_output(DIR_B)

# Spolecne smlouvy
all_bases = sorted(set(res_a.keys()) | set(res_b.keys()))
common = sorted(set(res_a.keys()) & set(res_b.keys()))

print(f'Smluv v {LABEL_A}: {len(res_a)}')
print(f'Smluv v {LABEL_B}: {len(res_b)}')
print(f'Spolecnych: {len(common)}\n')

# Srovnani po smlouvach
header = f'{"Smlouva":<20} {"PII-A":>5} {"PII-B":>5} {"Phan-A":>6} {"Phan-B":>6} {"TxJ-A":>5} {"TxJ-B":>5}  Lepsi'
print(header)
print('-' * len(header))

a_better = 0
b_better = 0
tie = 0

for base in common:
    ra = res_a[base]
    rb = res_b[base]
    pii_a = len(ra['pii_leaks'])
    pii_b = len(rb['pii_leaks'])
    ph_a = ra['phantom_tags']
    ph_b = rb['phantom_tags']
    tj_a = ra['txt_json_diff']
    tj_b = rb['txt_json_diff']

    score_a = pii_a * 3 + ph_a + tj_a
    score_b = pii_b * 3 + ph_b + tj_b

    if score_a < score_b:
        winner = f'<- {LABEL_A}'
        a_better += 1
    elif score_b < score_a:
        winner = f'-> {LABEL_B}'
        b_better += 1
    else:
        winner = '   shodne'
        tie += 1

    mark = ''
    if pii_a != pii_b or ph_a != ph_b or tj_a != tj_b:
        mark = ' ***'

    print(f'{base:<20} {pii_a:>5} {pii_b:>5} {ph_a:>6} {ph_b:>6} {tj_a:>5} {tj_b:>5}  {winner}{mark}')

print()
print(f'=== CELKOVE SROVNANI ({len(common)} spolecnych smluv) ===')
print(f'{LABEL_A} lepsi: {a_better}')
print(f'{LABEL_B} lepsi: {b_better}')
print(f'Shodne: {tie}')
print()

total_pii_a = sum(len(res_a[b]['pii_leaks']) for b in common)
total_pii_b = sum(len(res_b[b]['pii_leaks']) for b in common)
total_ph_a = sum(res_a[b]['phantom_tags'] for b in common)
total_ph_b = sum(res_b[b]['phantom_tags'] for b in common)
total_tj_a = sum(res_a[b]['txt_json_diff'] for b in common)
total_tj_b = sum(res_b[b]['txt_json_diff'] for b in common)

print(f'{"Metrika":<25} {LABEL_A:>12} {LABEL_B:>12}')
print(f'{"PII leaky celkem":<25} {total_pii_a:>12} {total_pii_b:>12}')
print(f'{"Phantom tagy celkem":<25} {total_ph_a:>12} {total_ph_b:>12}')
print(f'{"TXT/JSON rozdily celkem":<25} {total_tj_a:>12} {total_tj_b:>12}')
print()

# Detail PII leaku
for label, res, dirname in [(LABEL_A, res_a, DIR_A), (LABEL_B, res_b, DIR_B)]:
    leaks = []
    for b in common:
        for typ, orig in res[b]['pii_leaks']:
            leaks.append((b, typ, orig))
    if leaks:
        print(f'--- PII leaky v {label} ---')
        for b, typ, orig in leaks[:25]:
            print(f'  {b}: [{typ}] {orig}')
        if len(leaks) > 25:
            print(f'  ... +{len(leaks)-25} dalsich')
        print()
    else:
        print(f'--- PII leaky v {label}: ZADNE ---\n')
