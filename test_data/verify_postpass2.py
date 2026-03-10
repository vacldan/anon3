#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Extended check: look for inflected/possessive forms of first names in anon docs."""
import json, re, sys
from pathlib import Path
from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CASE_SUFFIXES = [
    "", "a", "u", "ovi", "em", "e", "ě", "y", "ou", "i",
    "ina", "iny", "ině", "inu", "inou",
    "ův", "ova", "ově", "ovu", "ovým",
]

IGNORE_WORDS = {
    'nová', 'nové', 'nového', 'nový', 'stav', 'stavu', 'stavě',
    'bílá', 'bílé', 'malá', 'malé', 'černá', 'černé',
    'město', 'místa', 'města', 'místo', 'soud', 'soudu',
    'svědci', 'svědek', 'rada', 'rady', 'pane', 'paní', 'pan',
    'české', 'česká', 'český', 'českou',
    'plné', 'plného', 'plná', 'plnou', 'celé', 'celou',
    'další', 'dalšího', 'jiné', 'jiného',
}

anon_dir = Path(__file__).parent / "anon_output"
results = []
total_checked = 0

for map_f in sorted(anon_dir.glob("*_map.json")):
    base = map_f.stem.replace("_map", "")
    anon_f = anon_dir / f"{base}_anon.docx"
    if not anon_f.exists():
        continue

    total_checked += 1
    data = json.loads(map_f.read_text(encoding="utf-8"))
    entities = data.get("entities", [])

    doc = Document(str(anon_f))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + "\n".join(p.text for p in cell.paragraphs)

    issues = []
    seen = set()
    for ent in entities:
        if ent.get("type") != "PERSON":
            continue
        orig = str(ent.get("original", "")).strip()
        if not orig or len(orig) < 5:
            continue
        parts = orig.split()
        if len(parts) < 2:
            continue
        first_name = parts[0]
        if len(first_name) < 4 or first_name.lower() in IGNORE_WORDS:
            continue

        stem = first_name.rstrip("a") if first_name.endswith("a") else first_name
        if len(stem) < 3:
            continue

        forms = set()
        for suf in CASE_SUFFIXES:
            form = stem + suf
            if len(form) >= 4 and form.lower() not in IGNORE_WORDS:
                forms.add(form)
        forms.add(first_name)

        for form in forms:
            if form.lower() in seen:
                continue
            seen.add(form.lower())
            for m in re.finditer(r"\b" + re.escape(form) + r"\b", text, re.IGNORECASE):
                before = text[max(0, m.start() - 5): m.start()]
                if "[[" not in before and "]]" not in before:
                    ctx = text[max(0, m.start() - 40): m.end() + 40].replace("\n", " ")
                    issues.append(f"'{m.group()}' (variant of '{first_name}' from '{orig}'): ...{ctx}...")

    if issues:
        results.append((base, issues))

print(f"Checked {total_checked} contracts")
print(f"Contracts with remaining inflected first-name leaks: {len(results)}")
for fname, issues in results:
    print(f"\n=== {fname} ===")
    for iss in issues[:8]:
        print(f"  {iss}")
    if len(issues) > 8:
        print(f"  ... and {len(issues) - 8} more")
