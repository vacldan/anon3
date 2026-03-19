import json, re, sys
from pathlib import Path
from docx import Document

TEST_DIR = Path(__file__).resolve().parent


def check(base_name: str, map_path: Path, anon_path: Path, src_path: Path):
    with open(map_path, encoding='utf-8') as f:
        mdata = json.load(f)

    src = Document(str(src_path))
    src_parts = [p.text for p in src.paragraphs]
    for t in src.tables:
        for row in t.rows:
            for cell in row.cells:
                src_parts.append(cell.text)
    src_text = '\n'.join(src_parts)

    anon = Document(str(anon_path))
    anon_parts = [p.text for p in anon.paragraphs]
    for t in anon.tables:
        for row in t.rows:
            for cell in row.cells:
                anon_parts.append(cell.text)
    anon_text = '\n'.join(anon_parts)
    clean = re.sub(r'\[\[\w+_\d+\]\]', '___', anon_text)

    persons = [e for e in mdata['entities'] if e['type'] == 'PERSON']
    phones = [e for e in mdata['entities'] if e['type'] == 'PHONE']
    insurances = [e for e in mdata['entities'] if e['type'] == 'INSURANCE_ID']
    plates = [e for e in mdata['entities'] if e['type'] == 'LICENSE_PLATE']

    print(f'\n====== {base_name} ======')
    print(f'Persons: {len(persons)} | Phones: {len(phones)} | Insurance: {len(insurances)} | Plates: {len(plates)}')

    # Print persons
    seen = set()
    for p in persons:
        if p['label'] not in seen:
            seen.add(p['label'])
            variants = [e['original'] for e in persons if e['label'] == p['label'] and e['original'] != p['original']]
            vstr = ' - ' + ', '.join(variants) if variants else ''
            print(f"  {p['label']}: {p['original']}{vstr}")

    # Check gender mismatches
    for p in persons:
        if p['label'] in seen:
            parts = p['original'].strip().split()
            if len(parts) == 2:
                first, last = parts
                first_lo = first.lower()
                last_lo = last.lower()
                if not first_lo.endswith('a') and last_lo.endswith(('ová', 'á')) and not last_lo.endswith('ská'):
                    print(f'  ** GENDER MISMATCH: {p["label"]}: {p["original"]}')

    # Check name leaks
    names_to_check = set()
    CZ_UPPER = r'A-ZÁČĎÉĚÍŇÓŘŠŤÚŮÝŽ'
    CZ_LOWER = r'a-záčďéěíňóřšťúůýž'
    for m in re.finditer(rf'([{CZ_UPPER}][{CZ_LOWER}]+)\s+([{CZ_UPPER}][{CZ_LOWER}]+)', src_text):
        first, last = m.group(1), m.group(2)
        skip = {'Tento', 'Další', 'Hlavní', 'Celkem', 'Dokument', 'Datum',
                'Komplexní', 'Rozsáhlé', 'Bankovní', 'Firemní', 'Dětská',
                'Pacient', 'Pacientka', 'Klient', 'Dlužník', 'Matkou',
                'Svědkem', 'Obžalovaný', 'Poškozený', 'Odsouzený',
                'Obviněný', 'Rozvedený', 'Účastník', 'Správce',
                'Vedoucí', 'Sociální', 'Okresní', 'Městský', 'Obvodní',
                'Krajský', 'Nejvyšší', 'Ústavní', 'Speciální', 'Civilní',
                'Trestní', 'Vzácná', 'Pěstouni'}
        if first in skip or len(first) < 3 or len(last) < 3:
            continue
        names_to_check.add(f'{first} {last}')

    leaks = []
    for name in sorted(names_to_check):
        if name in clean:
            idx = clean.find(name)
            ctx = clean[max(0,idx-20):idx+len(name)+20]
            leaks.append(f'{name}  | ...{ctx}...')

    if leaks:
        print(f'  NAME LEAKS ({len(leaks)}):')
        for l in leaks:
            print(f'    - {l}')
    else:
        print('  No name leaks.')

    # Check pojistenec as phone
    pojist_nums = set()
    for m in re.finditer(r'(?:pojišt[eě]n[ecka]\w*|VZP|OZP|ZP|ČPZP)\s*[:\s]\s*(\d{9,10})', src_text, re.IGNORECASE):
        pojist_nums.add(m.group(1))
    bad_phones = []
    for ph in phones:
        clean_ph = re.sub(r'[\s+]', '', ph['original'])
        if clean_ph in pojist_nums:
            bad_phones.append(f"{ph['label']}: {ph['original']}")
    if bad_phones:
        print(f'  POJISTENEC AS PHONE ({len(bad_phones)}):')
        for bp in bad_phones:
            print(f'    - {bp}')

    # Check medical protocols as plates
    for pl in plates:
        val = pl['original']
        if re.match(r'^[A-Z]{2}\d{4}$', val):
            idx = src_text.find(val)
            if idx >= 0:
                ctx = src_text[max(0,idx-80):idx+len(val)+40].lower()
                if any(w in ctx for w in ['protokol', 'chemoterapie', 'leukemi', 'neuroblast', 'onkolog']):
                    print(f'  FALSE PLATE: {pl["label"]}: {val} (medical protocol)')

    # Check incomplete persons (only surname, no first name)
    for p in persons:
        if p['label'] in seen:
            parts = p['original'].strip().split()
            if len(parts) == 1:
                print(f'  INCOMPLETE: {p["label"]}: "{p["original"]}" (surname only)')

# Najdi všechny páry (anon, map, src) – všechny smlouvy
anon_files = sorted(TEST_DIR.glob("*_anon.docx"))
for anon_path in anon_files:
    base = anon_path.stem.replace("_anon", "")
    if "_deanon" in base:
        continue
    map_path = TEST_DIR / f"{base}_map.json"
    src_path = TEST_DIR / f"{base}.docx"
    if not map_path.exists() or not src_path.exists():
        continue
    check(base, map_path, anon_path, src_path)
