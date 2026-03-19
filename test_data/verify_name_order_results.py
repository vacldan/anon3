#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ověří výsledky anonymizace 10 smluv name_order:
1. Mapa obsahuje všechna očekávaná jména
2. V anonymizovaném dokumentu nejsou viditelná původní jména (leaky)
"""

import json
import re
from pathlib import Path

from docx import Document

TEST_DIR = Path(__file__).resolve().parent

# Očekávaná jména v každé smlouvě (nominativ, canonical)
EXPECTED_NAMES = {
    "01": [("Jan", "Novák"), ("Jana", "Nováková")],
    "02": [("Jan", "Novák"), ("Jana", "Nováková")],
    "03": [("Jan", "Novák"), ("Jana", "Nováková")],
    "04": [("Pavel", "Dvořák")],
    "05": [("Tomáš", "Svoboda")],
    "06": [("Jan", "Novák"), ("Jana", "Nováková")],
    "07": [("Tomáš", "Svoboda")],
    "08": [("Petra", "Dvořáková"), ("Pavel", "Dvořák")],
    "09": [("Jan", "Novák"), ("Pavel", "Novák"), ("Jana", "Svobodová"), ("Tomáš", "Dvořák"), ("Petra", "Dvořáková")],
    "10": [("Marek", "Procházka")],
}


def get_text_from_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    return " ".join(parts)


def get_persons_from_map(map_path: Path) -> list[dict]:
    with open(map_path, encoding="utf-8") as f:
        data = json.load(f)
    return [e for e in data.get("entities", []) if e.get("type") == "PERSON"]


def check_leaks(anon_text: str, origins: list[str]) -> list[str]:
    """Vrátí seznam originálů, které se objevují v anonymizovaném textu mimo tag."""
    leaks = []
    clean = re.sub(r"\[\[[^\]]+\]\]", " ", anon_text)
    for orig in origins:
        if len(orig) < 3:
            continue
        if orig.lower() in clean.lower():
            leaks.append(orig)
    return leaks


def main():
    if hasattr(__import__("sys").stdout, "reconfigure"):
        __import__("sys").stdout.reconfigure(encoding="utf-8", errors="replace")
    print("=" * 70)
    print("VERIFIKACE ANONYMIZACE - VŠECHNY SMLOUVY (osoby + leaky)")
    print("=" * 70)

    pairs = []
    for anon_path in sorted(TEST_DIR.glob("*_anon.docx")):
        stem = anon_path.stem.replace("_anon", "")
        if "_deanon" in stem:
            continue
        map_path = TEST_DIR / f"{stem}_map.json"
        if map_path.exists():
            pairs.append((stem, map_path, anon_path))

    print(f"Nalezeno {len(pairs)} párů (anon + mapa)\n")

    all_ok = True
    for stem, map_path, anon_path in pairs:
        persons = get_persons_from_map(map_path)
        anon_text = get_text_from_docx(anon_path)

        # 1. Počet unikátních osob (podle labelu PERSON_1, PERSON_2, ...)
        unique_labels = set(e.get("label", "") for e in persons if e.get("label"))
        person_count = len(unique_labels)

        # 2. Všechny originály pro leak check
        all_origins = list(set(str(e.get("original", "")).strip() for e in persons if e.get("original")))
        leaks = check_leaks(anon_text, all_origins)

        # 3. Ukázka originálů z mapy
        sample_origins = list(all_origins)[:4]

        expected = EXPECTED_NAMES.get(stem.split("_")[-1], []) if "name_order" in stem else []
        expected_count = len(set((f, l) for f, l in expected)) if expected else 0
        status = "OK" if not leaks and (expected_count == 0 or person_count >= expected_count) else "PROBLEM"
        if leaks or (expected_count > 0 and person_count < expected_count):
            all_ok = False

        exp_str = f" (ocekavano min. {expected_count})" if expected_count else ""
        print(f"\n[{stem}] {status}")
        print(f"  Osob v mape: {person_count}{exp_str}")
        print(f"  Ukazka originalu: {sample_origins}")
        if leaks:
            print(f"  LEAKY (original v textu): {leaks}")

    print("\n" + "=" * 70)
    print("VÝSLEDEK: " + ("VŠE OK" if all_ok else "NALEZENY PROBLÉMY"))
    print("=" * 70)


if __name__ == "__main__":
    main()
