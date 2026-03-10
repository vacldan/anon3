#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify no district names leak after ADDRESS tags (e.g. [[ADDRESS_2]] - Vinohrady)."""
import re, sys
from pathlib import Path
from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

KNOWN_DISTRICTS = [
    "Vinohrady", "Smíchov", "Žižkov", "Karlín", "Dejvice", "Holešovice",
    "Vršovice", "Nusle", "Košíře", "Břevnov", "Bubeneč", "Libeň",
    "Podolí", "Braník", "Krč", "Modřany", "Chodov", "Háje",
    "Prosek", "Letňany", "Černý Most", "Střížkov", "Bohnice",
    "Troja", "Vysočany", "Hloubětín", "Hostivař", "Záběhlice",
    "Michle", "Kunratice", "Řepy", "Zličín", "Stodůlky",
    "Staré Město", "Nové Město", "Malá Strana", "Hradčany",
    "Střed", "Sever", "Jih", "Bystrc", "Líšeň", "Královo Pole",
    "Poruba", "Dubina", "Hrabůvka", "Jižní Předměstí",
]

anon_dir = Path(__file__).parent / "anon_output"
results = []
total_checked = 0

for anon_f in sorted(anon_dir.glob("*_anon.docx")):
    total_checked += 1
    doc = Document(str(anon_f))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + "\n".join(p.text for p in cell.paragraphs)

    issues = []

    # Check 1: [[ADDRESS_N]] - DistrictName (should have been absorbed)
    for m in re.finditer(r'\[\[ADDRESS_\d+\]\]\s*-\s*([A-ZÁČĎÉĚÍŇÓŘŠŤÚŮÝŽ][a-záčďéěíňóřšťúůýž]+)', text):
        district = m.group(1)
        ctx = text[max(0, m.start() - 20): m.end() + 20].replace("\n", " ")
        issues.append(f"DISTRICT AFTER TAG: '{district}' -> ...{ctx}...")

    # Check 2: Known district names outside tags
    for d in KNOWN_DISTRICTS:
        for m in re.finditer(r'\b' + re.escape(d) + r'\b', text):
            before = text[max(0, m.start() - 5): m.start()]
            if "[[" not in before and "]]" not in before:
                ctx = text[max(0, m.start() - 40): m.end() + 40].replace("\n", " ")
                issues.append(f"DISTRICT VISIBLE: '{d}' -> ...{ctx}...")

    if issues:
        results.append((anon_f.stem, issues))

print(f"Checked {total_checked} contracts")
print(f"Contracts with district leaks: {len(results)}")
for fname, issues in results:
    print(f"\n=== {fname} ===")
    for iss in issues[:5]:
        print(f"  {iss}")
    if len(issues) > 5:
        print(f"  ... and {len(issues) - 5} more")
