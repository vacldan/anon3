# -*- coding: utf-8 -*-
"""
Porovná dva DOCX dokumenty
"""

import sys
from docx import Document

def extract_text(docx_path):
    """Extrahuje text z DOCX dokumentu"""
    doc = Document(docx_path)
    text_parts = []

    # Hlavní text
    for para in doc.paragraphs:
        text_parts.append(para.text)

    # Text z tabulek
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text_parts.append(para.text)

    return '\n'.join(text_parts)

def compare_documents(doc1_path, doc2_path):
    """Porovná dva dokumenty"""
    print(f"Porovnavam:")
    print(f"  {doc1_path}")
    print(f"  {doc2_path}")
    print()

    text1 = extract_text(doc1_path)
    text2 = extract_text(doc2_path)

    # Porovnej délky
    print(f"Delka dokumentu 1: {len(text1)} znaku")
    print(f"Delka dokumentu 2: {len(text2)} znaku")
    print()

    # Zkontroluj, jestli jsou identické
    if text1 == text2:
        print("✓ DOKUMENTY JSOU IDENTICKÉ!")
        return True
    else:
        print("✗ Dokumenty se liší")

        # Najdi rozdíly
        lines1 = text1.split('\n')
        lines2 = text2.split('\n')

        print(f"\nPocet radku dokument 1: {len(lines1)}")
        print(f"Pocet radku dokument 2: {len(lines2)}")
        print()

        # Najdi první 10 rozdílů
        diff_count = 0
        for i, (line1, line2) in enumerate(zip(lines1, lines2)):
            if line1 != line2 and diff_count < 10:
                print(f"Rozdil na radku {i+1}:")
                print(f"  Dokument 1: {line1[:100]}")
                print(f"  Dokument 2: {line2[:100]}")
                print()
                diff_count += 1

        if diff_count >= 10:
            print("... a dalsi rozdily")

        return False

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Pouziti: python compare_documents.py dokument1.docx dokument2.docx")
        sys.exit(1)

    doc1 = sys.argv[1]
    doc2 = sys.argv[2]

    result = compare_documents(doc1, doc2)
    sys.exit(0 if result else 1)
