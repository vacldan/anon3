# pdf2docx_cli.py - PDF to DOCX converter with OCR fallback for scanned documents
import sys
import os
import re
import subprocess
from pathlib import Path

# Set UTF-8 encoding - safe cross-platform approach
if sys.platform == 'win32':
    import io
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')


# --- Auto-detekce Tesseract a Poppler na Windows ---

def _setup_windows_paths():
    """Najde Tesseract a Poppler na běžných Windows cestách a přidá do PATH."""
    if sys.platform != 'win32':
        return

    # Běžné cesty k Tesseract
    tesseract_dirs = [
        r"C:\Program Files\Tesseract-OCR",
        r"C:\Program Files (x86)\Tesseract-OCR",
        r"C:\Tesseract-OCR",
    ]
    for d in tesseract_dirs:
        exe = os.path.join(d, "tesseract.exe")
        if os.path.isfile(exe):
            if d not in os.environ.get("PATH", ""):
                os.environ["PATH"] = d + ";" + os.environ.get("PATH", "")
            break

    # Běžné cesty k Poppler (pdftoppm, pdftotext)
    poppler_dirs = []
    for drive in ["C:"]:
        base = drive + "\\"
        try:
            for name in os.listdir(base):
                if name.lower().startswith("poppler"):
                    lib_bin = os.path.join(base, name, "Library", "bin")
                    if os.path.isdir(lib_bin):
                        poppler_dirs.append(lib_bin)
                    plain_bin = os.path.join(base, name, "bin")
                    if os.path.isdir(plain_bin):
                        poppler_dirs.append(plain_bin)
        except OSError:
            pass
    # Také zkontroluj Program Files
    for pf in [r"C:\Program Files", r"C:\Program Files (x86)"]:
        try:
            for name in os.listdir(pf):
                if name.lower().startswith("poppler"):
                    lib_bin = os.path.join(pf, name, "Library", "bin")
                    if os.path.isdir(lib_bin):
                        poppler_dirs.append(lib_bin)
                    plain_bin = os.path.join(pf, name, "bin")
                    if os.path.isdir(plain_bin):
                        poppler_dirs.append(plain_bin)
        except OSError:
            pass

    for d in poppler_dirs:
        if d not in os.environ.get("PATH", ""):
            os.environ["PATH"] = d + ";" + os.environ.get("PATH", "")
            break


_setup_windows_paths()

# --- Kontrola závislostí ---

_has_pdf2docx = False
try:
    from pdf2docx import Converter
    _has_pdf2docx = True
except ImportError:
    pass

_has_ocr = False
try:
    import pytesseract
    from pdf2image import convert_from_path
    from PIL import Image, ImageEnhance, ImageFilter
    from docx import Document
    from docx.shared import Pt, Cm
    _has_ocr = True

    # Explicitně nastav cestu k tesseract.exe na Windows
    if sys.platform == 'win32':
        for d in [r"C:\Program Files\Tesseract-OCR",
                   r"C:\Program Files (x86)\Tesseract-OCR",
                   r"C:\Tesseract-OCR"]:
            exe = os.path.join(d, "tesseract.exe")
            if os.path.isfile(exe):
                pytesseract.pytesseract.tesseract_cmd = exe
                break
except ImportError:
    pass

if not _has_pdf2docx and not _has_ocr:
    print("ERROR: Zadna knihovna pro konverzi PDF neni nainstalována!", flush=True)
    print("Pro textove PDF:    pip install pdf2docx", flush=True)
    print("Pro skenovane PDF:  pip install pytesseract pdf2image Pillow python-docx", flush=True)
    print("                    apt install tesseract-ocr tesseract-ocr-ces poppler-utils", flush=True)
    sys.exit(1)


# --- Detekce typu PDF ---

def is_scanned_pdf(pdf_path: Path) -> bool:
    """
    Zjistí, zda je PDF skenované (obrázkové) nebo textové.

    Textové PDF → pdftotext vytáhne text, není potřeba OCR.
    Skenované PDF → pdftotext nevytáhne skoro nic, potřeba OCR.
    """
    try:
        result = subprocess.run(
            ['pdftotext', str(pdf_path), '-'],
            capture_output=True, timeout=30,
            encoding='utf-8', errors='replace'
        )
        text = (result.stdout or '').strip()
        # Méně než 50 znaků na celý dokument = pravděpodobně sken
        return len(text) < 50
    except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
        # pdftotext není dostupný nebo timeout → zkus OCR pro jistotu
        return True


# --- Konverze textového PDF (pdf2docx) ---

def convert_text_pdf(pdf_path: Path, docx_path: Path) -> bool:
    """Převede textové PDF do DOCX přes knihovnu pdf2docx."""
    if not _has_pdf2docx:
        return False

    try:
        print("  Metoda: pdf2docx (textove PDF)", flush=True)
        cv = Converter(str(pdf_path))
        cv.convert(str(docx_path), start=0, end=None)
        cv.close()

        if docx_path.exists() and docx_path.stat().st_size > 0:
            return True
        return False

    except Exception as e:
        print(f"  pdf2docx selhalo: {e}", flush=True)
        return False


# --- Konverze skenovaného PDF (OCR) ---

def preprocess_image(image):
    """Předzpracuje obrázek pro lepší OCR kvalitu.

    Minimální zásah - jen převod na šedou a jemné zvýšení kontrastu.
    Upscale ani binarizace se nedělá (amplifikuje šum / ničí diacritiku).
    """
    # Převeď na šedou
    img = image.convert('L')

    # Jemné zvýšení kontrastu
    enhancer = ImageEnhance.Contrast(img)
    img = enhancer.enhance(1.3)

    return img


def cleanup_ocr_text(text: str) -> str:
    """Vyčistí běžné OCR artefakty."""
    lines = text.split('\n')
    cleaned = []
    for line in lines:
        stripped = line.strip()
        # Odstraň řádky s jen symboly (artefakty skenování)
        if stripped and re.match(r'^[|_\-=~*#@!—–―]+$', stripped):
            continue
        # Odstraň řádky kde většina znaků jsou speciální/nesmyslné (artefakty loga, čárových kódů)
        if stripped and len(stripped) > 3:
            alnum = sum(1 for c in stripped if c.isalnum() or c in ' .,;:')
            if alnum < len(stripped) * 0.4:
                continue
        # Odstraň zápatí s kódy dokumentu (HYVS01. 20190322 cs41000...)
        if stripped and re.match(r'^[A-Z]{2,6}\d{2,4}[\.\s]', stripped) and len(stripped) < 60:
            if sum(1 for c in stripped if c.isdigit()) > len(stripped) * 0.3:
                continue
        # Odstraň nesmyslné řádky z hlavičky/loga (hodně krátkých uppercase slov za sebou)
        if stripped and len(stripped) > 10:
            words = stripped.split()
            if len(words) >= 4:
                short_upper = sum(1 for w in words if w.isupper() and len(w) <= 4 and not re.match(r'^\d', w))
                if short_upper >= len(words) * 0.6:
                    continue
        # Odstraň krátké artefakty (1-3 znaky, které nejsou čísla oddílů ani písmena v závorkách)
        if stripped and len(stripped) <= 3:
            if not re.match(r'^\d+[\.\)]?$', stripped) and not re.match(r'^\([a-z]\)$', stripped) and stripped != 'a':
                continue
        cleaned.append(line)

    text = '\n'.join(cleaned)

    # Odstraň | na začátku řádků (OCR artefakt svislých čar)
    text = re.sub(r'^\|\s*', '', text, flags=re.MULTILINE)

    # Oprav dvojité mezery
    text = re.sub(r'  +', ' ', text)

    # Oprav rozlomená slova na konci řádku (čes-\nký → český)
    text = re.sub(r'(\w)-\n(\w)', r'\1\2', text)

    # Sloučí osamocené číslo bodu (např. "2.2\n\nJakou nemovitou...")
    # s následujícím řádkem POUZE pokud ten začíná písmenem (ne dalším číslem)
    text = re.sub(
        r'^(\d{1,3}[\.\)]{0,1}\d{0,2}[\'\'"]?)\s*\n+(?=[A-Za-z\u00C0-\u024F])',
        r'\1 ',
        text,
        flags=re.MULTILINE
    )

    return text


def _merge_lines_to_paragraphs(text: str) -> list:
    """Sloučí OCR řádky do logických odstavců.

    OCR produkuje řádky podle vizuálního zlomu na stránce.
    Tato funkce je sloučí do odstavců - nový odstavec začíná při:
    prázdném řádku, čísle oddílu, (a)/(b), odrážce -.
    Pokud text plynule pokračuje přes prázdný řádek (předchozí nekončí
    tečkou a další začíná malým písmenem), sloučí se.
    """
    lines = text.split('\n')
    paragraphs = []
    current = []

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            # Prázdný řádek - podívej se jestli text pokračuje plynule
            # (předchozí nekončí tečkou/dvojtečkou a další začíná malým písmenem)
            if current:
                # Najdi další neprázdný řádek
                next_idx = i + 1
                while next_idx < len(lines) and not lines[next_idx].strip():
                    next_idx += 1

                if next_idx < len(lines):
                    next_stripped = lines[next_idx].strip()
                    prev_text = current[-1] if current else ''
                    # Pokud předchozí text nekončí větnou interpunkcí a další
                    # začíná malým písmenem nebo spojkou → sloučit
                    continues = (
                        prev_text and
                        next_stripped and
                        (
                            # Předchozí nekončí větnou interpunkcí, další začíná malým písmenem
                            (not prev_text[-1] in '.!?:' and
                             (next_stripped[0].islower() or next_stripped in ('a', 'i', 'nebo')))
                            or
                            # Další začíná pokračovací interpunkcí (: , ;)
                            next_stripped[0] in ':,;'
                        )
                    )
                    if continues:
                        i += 1
                        continue

                paragraphs.append(' '.join(current))
                current = []
            paragraphs.append('')  # zachovej prázdný řádek
            i += 1
            continue

        # Nový odstavec začíná při: číslo oddílu, (a)/(b), odrážka -, nadpis velkými
        is_new_para = (
            re.match(r'^\d{1,2}[\.\)]\s*\d{0,2}', stripped) or  # 2.1, 3., 4)
            re.match(r'^\([a-z]\)', stripped) or                  # (a), (b)
            re.match(r'^-\s', stripped) or                        # - odrážka
            re.match(r'^[A-ZÁČĎÉĚÍŇÓŘŠŤÚŮÝŽ]{3,}', stripped)    # NADPIS
        )

        if is_new_para and current:
            paragraphs.append(' '.join(current))
            current = []

        current.append(stripped)
        i += 1

    if current:
        paragraphs.append(' '.join(current))

    return paragraphs


def convert_scanned_pdf(pdf_path: Path, docx_path: Path,
                        lang: str = "ces", dpi: int = 300, psm: int = 6) -> bool:
    """Převede skenované PDF do DOCX přes Tesseract OCR."""
    if not _has_ocr:
        print("  OCR neni dostupne! Nainstalujte:", flush=True)
        print("    pip install pytesseract pdf2image Pillow python-docx", flush=True)
        print("    apt install tesseract-ocr tesseract-ocr-ces poppler-utils", flush=True)
        return False

    try:
        print(f"  Metoda: Tesseract OCR (skenovane PDF)", flush=True)
        print(f"  Jazyk: {lang}, Rozliseni: {dpi} DPI, PSM: {psm}", flush=True)

        # 1. PDF → obrázky
        print("  Prevadim stranky na obrazky...", flush=True)
        images = convert_from_path(str(pdf_path), dpi=dpi)
        print(f"  Nalezeno {len(images)} stranek", flush=True)

        # 2. OCR každé stránky → text, sloučit do celku
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        all_text_parts = []

        for i, image in enumerate(images, 1):
            pct = int(i / len(images) * 100)
            print(f"\r  OCR: {i}/{len(images)} stranek ({pct}%)", end='', flush=True)

            # Předzpracování obrázku pro lepší kvalitu OCR
            processed = preprocess_image(image)

            page_text = pytesseract.image_to_string(
                processed, lang=lang, config=f"--oem 1 --psm {psm}"
            )
            page_text = cleanup_ocr_text(page_text)
            all_text_parts.append(page_text)

        print(flush=True)

        # Sloučit text ze všech stránek do jednoho celku
        full_text = '\n'.join(all_text_parts)
        total_chars = len(full_text)

        # Odstraň nadbytečné prázdné řádky (3+ za sebou → max 2)
        full_text = re.sub(r'\n{3,}', '\n\n', full_text)

        # 3. Text → DOCX odstavce (sloučené řádky + odsazení)
        merged_paragraphs = _merge_lines_to_paragraphs(full_text)
        for para_text in merged_paragraphs:
            text = para_text.strip()
            para = doc.add_paragraph(text if text else '')
            if text:
                # Odsazení pro (a), (b), (c)... a odrážky -
                if re.match(r'^\([a-z]\)', text):
                    para.paragraph_format.left_indent = Cm(1.0)
                elif re.match(r'^-\s', text):
                    para.paragraph_format.left_indent = Cm(1.5)

        doc.save(str(docx_path))
        print(f"  OCR hotovo, rozpoznano {total_chars:,} znaku", flush=True)

        return docx_path.exists() and docx_path.stat().st_size > 0

    except Exception as e:
        print(f"  OCR selhalo: {e}", flush=True)
        return False


# --- Konverze obrázku (PNG/JPG/TIFF/BMP) přes OCR ---

IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.tiff', '.tif', '.bmp', '.webp'}


def convert_image(image_path: Path, lang: str = "ces", psm: int = 6) -> bool:
    """Převede obrázek (PNG/JPG/TIFF/BMP) do DOCX přes Tesseract OCR."""
    if not _has_ocr:
        print("  OCR neni dostupne! Nainstalujte:", flush=True)
        print("    pip install pytesseract Pillow python-docx", flush=True)
        print("    apt install tesseract-ocr tesseract-ocr-ces", flush=True)
        return False

    docx_path = image_path.with_suffix('.docx')

    try:
        print(f"  Metoda: Tesseract OCR (obrazek)", flush=True)
        print(f"  Jazyk: {lang}, PSM: {psm}", flush=True)

        # 1. Načti obrázek
        image = Image.open(str(image_path))
        print(f"  Obrazek: {image.size[0]}x{image.size[1]} px", flush=True)

        # 2. Předzpracování
        processed = preprocess_image(image)

        # 3. OCR
        print("  Spoustim OCR...", flush=True)
        page_text = pytesseract.image_to_string(
            processed, lang=lang, config=f"--oem 1 --psm {psm}"
        )
        page_text = cleanup_ocr_text(page_text)
        total_chars = len(page_text)

        # Odstraň nadbytečné prázdné řádky
        page_text = re.sub(r'\n{3,}', '\n\n', page_text)

        # 4. Text → DOCX
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)

        merged_paragraphs = _merge_lines_to_paragraphs(page_text)
        for para_text in merged_paragraphs:
            text = para_text.strip()
            para = doc.add_paragraph(text if text else '')
            if text:
                if re.match(r'^\([a-z]\)', text):
                    para.paragraph_format.left_indent = Cm(1.0)
                elif re.match(r'^-\s', text):
                    para.paragraph_format.left_indent = Cm(1.5)

        doc.save(str(docx_path))
        print(f"  OCR hotovo, rozpoznano {total_chars:,} znaku", flush=True)

        return docx_path.exists() and docx_path.stat().st_size > 0

    except Exception as e:
        print(f"  OCR obrazku selhalo: {e}", flush=True)
        return False


# --- Hlavní konverzní funkce ---

def convert_file(file_path: Path, dpi: int = 300, lang: str = "ces", psm: int = 6) -> bool:
    """
    Převede PDF nebo obrázek do DOCX - automaticky zvolí správnou metodu.

    Textové PDF  → pdf2docx (zachová formátování)
    Skenované PDF → Tesseract OCR (přečte text z obrázku)
    Obrázek (PNG/JPG/TIFF/BMP) → Tesseract OCR
    """
    docx_path = file_path.with_suffix('.docx')

    print(f"Zpracovavam: {file_path.name}", flush=True)
    print(f"Vystupni DOCX: {docx_path.name}", flush=True)

    if not file_path.exists():
        print(f"ERROR: Soubor neexistuje: {file_path}", flush=True)
        return False

    if file_path.stat().st_size == 0:
        print(f"ERROR: Soubor je prazdny: {file_path}", flush=True)
        return False

    # Obrázek → OCR přímo
    if file_path.suffix.lower() in IMAGE_EXTENSIONS:
        print("  Detekovano: OBRAZEK", flush=True)
        success = convert_image(file_path, lang=lang, psm=psm)
    else:
        # PDF → detekce textové/skenované
        scanned = is_scanned_pdf(file_path)

        if scanned:
            print("  Detekovano: SKENOVANE PDF (obrazek)", flush=True)
            success = convert_scanned_pdf(file_path, docx_path, lang=lang, dpi=dpi, psm=psm)
        else:
            print("  Detekovano: TEXTOVE PDF", flush=True)
            success = convert_text_pdf(file_path, docx_path)

            # Fallback: pokud pdf2docx selhalo a máme OCR, zkus OCR
            if not success and _has_ocr:
                print("  Zkousim fallback pres OCR...", flush=True)
                success = convert_scanned_pdf(file_path, docx_path, lang=lang, dpi=dpi, psm=psm)

    if success:
        print(f"[OK] Uspesne prevedeno: {docx_path.name}", flush=True)
        print(f"Velikost: {docx_path.stat().st_size} bytu", flush=True)
    else:
        print("ERROR: Konverze selhala", flush=True)

    return success


# Zpětná kompatibilita
def convert_pdf(pdf_path: Path, dpi: int = 300, lang: str = "ces", psm: int = 6) -> bool:
    return convert_file(pdf_path, dpi=dpi, lang=lang, psm=psm)


def main():
    print("PDF to DOCX Converter - Nixminds Document Suite", flush=True)
    print("=" * 50, flush=True)

    # Informace o dostupných metodách
    if _has_pdf2docx:
        print("  [+] pdf2docx: dostupne (textove PDF)", flush=True)
    else:
        print("  [-] pdf2docx: neni (pip install pdf2docx)", flush=True)

    if _has_ocr:
        print("  [+] Tesseract OCR: dostupne (skenovane PDF)", flush=True)
    else:
        print("  [-] Tesseract OCR: neni (pip install pytesseract pdf2image Pillow)", flush=True)

    # Parsuj argumenty - podpora --dpi, --lang, --psm
    args = sys.argv[1:]
    dpi = 300
    lang = "ces"
    psm = 6
    pdf_files = []

    i = 0
    while i < len(args):
        if args[i] == '--dpi' and i + 1 < len(args):
            dpi = int(args[i + 1])
            i += 2
        elif args[i] == '--lang' and i + 1 < len(args):
            lang = args[i + 1]
            i += 2
        elif args[i] == '--psm' and i + 1 < len(args):
            psm = int(args[i + 1])
            i += 2
        else:
            pdf_files.append(args[i])
            i += 1

    SUPPORTED_EXTENSIONS = {'.pdf'} | IMAGE_EXTENSIONS

    if not pdf_files:
        print("\nERROR: Nebyl zadan soubor", flush=True)
        print("Pouziti: python pdf2docx_cli.py <soubor> [--dpi 400] [--lang ces] [--psm 6]", flush=True)
        print("Podporovane formaty: PDF, PNG, JPG, JPEG, TIFF, TIF, BMP, WEBP", flush=True)
        sys.exit(1)

    success_count = 0
    total_count = 0

    for pdf_arg in pdf_files:
        total_count += 1
        file_path = Path(pdf_arg)

        print(f"\nZpracovavam ({total_count}): {file_path.name}", flush=True)

        if not file_path.exists():
            print(f"[!] Soubor neexistuje: {file_path}", flush=True)
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            print(f"[!] Nepodporovany format: {file_path.suffix}", flush=True)
            print(f"    Podporovane: {', '.join(sorted(SUPPORTED_EXTENSIONS))}", flush=True)
            continue

        if convert_file(file_path, dpi=dpi, lang=lang, psm=psm):
            success_count += 1
        else:
            print(f"[X] Konverze selhala pro: {file_path.name}", flush=True)

    print("\n" + "=" * 50, flush=True)
    print(f"VYSLEDEK: {success_count}/{total_count} souboru uspesne prevedeno", flush=True)

    exit_code = 0
    if success_count == total_count and total_count > 0:
        print("[OK] Vsechny soubory byly uspesne prevedeny!", flush=True)
    else:
        print("[!] Nektere soubory nebyly prevedeny", flush=True)
        exit_code = 1

    if not os.environ.get('NO_PAUSE'):
        try:
            input("\nStiskni Enter pro ukonceni...")
        except (EOFError, KeyboardInterrupt):
            pass

    sys.exit(exit_code)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback
        print(f"\nERROR: Neocekavana chyba: {e}", flush=True)
        print(traceback.format_exc(), flush=True)
        sys.exit(1)
