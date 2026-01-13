# deanonymizer.py - Reverse anonymization

import json, sys, os
from docx import Document

# Set UTF-8 encoding for Windows
import io
sys.stdout = io.TextIOWrapper(sys.stdout.detach(), encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.detach(), encoding='utf-8')

# Default file names (can be overridden by command line arguments)
if len(sys.argv) >= 4:
    IN_DOC = sys.argv[1]    # anonymized document
    IN_MAP = sys.argv[2]    # JSON mapping file
    OUT_DOC = sys.argv[3]   # output deanonymized document
else:
    IN_DOC = "smlouva_anon.docx"
    IN_MAP = "smlouva_map.json"
    OUT_DOC = "smlouva_deanon.docx"

print("DEANONYMIZER - obnoveni puvodniho dokumentu...")
print(f"Vstupni anonymni dokument: {IN_DOC}")
print(f"JSON mapa: {IN_MAP}")
print(f"Vystupni dokument: {OUT_DOC}")

def deanonymize_document(anon_doc_path, map_path, output_path):
    """Deanonymizuje dokument podle mapy"""

    # Načti mapu anonymizace
    try:
        with open(map_path, "r", encoding="utf-8") as f:
            map_data = json.load(f)

        # Převeď nový formát mapy (s entities) na jednoduchý slovník
        mapping = {}
        if isinstance(map_data, dict) and "entities" in map_data:
            # Nový formát s entities
            # Pro každý tag shromáždi všechny varianty
            tag_variants = {}
            for entity in map_data["entities"]:
                label = entity["label"]
                original = entity["original"]
                entity_type = entity.get("type", "")

                if label not in tag_variants:
                    tag_variants[label] = {"type": entity_type, "variants": []}
                tag_variants[label]["variants"].append(original)

            # Pro každý tag vyberi kanonickou formu
            for label, data in tag_variants.items():
                variants = data["variants"]
                entity_type = data["type"]

                if entity_type == "PERSON":
                    # Pro osoby: vyberi celé jméno (obsahuje mezeru) v nominativu
                    full_names = [v for v in variants if ' ' in v]
                    if full_names:
                        # Preferuj nominativ: křestní jméno NEKONČÍ na "-a" (pro mužská jména)
                        # nebo příjmení NEKONČÍ na "-ové/-ého" (pro genitiv)
                        nominative_candidates = []
                        for name in full_names:
                            parts = name.split()
                            if len(parts) >= 2:
                                first_name = parts[0]
                                last_name = parts[-1]
                                # Heuristika: nominativ mužského jména nekončí na "a",
                                # ženského jména ano, takže kontrolujeme jen koncovky příjmení
                                # Genitiv mužských příjmení: -a (Nováka), -ého (Novotného)
                                # Dativ: -ovi (Novákovi), -ému (Novotnému)
                                is_likely_nominative = not (
                                    last_name.endswith('ovi') or
                                    last_name.endswith('ému') or
                                    last_name.endswith('ého') or
                                    (last_name.endswith('a') and first_name.endswith('a'))  # Oba končí na -a = genitiv
                                )
                                if is_likely_nominative:
                                    nominative_candidates.append(name)

                        # Vybrat nejdelší nominativní tvar, nebo nejdelší celkový pokud není nominativ
                        if nominative_candidates:
                            mapping[label] = max(nominative_candidates, key=len)
                        else:
                            mapping[label] = max(full_names, key=len)
                    else:
                        # Pokud není celé jméno, vybrat nejdelší variantu
                        mapping[label] = max(variants, key=len)
                else:
                    # Pro ostatní entity: vybrat nejdelší variantu
                    mapping[label] = max(variants, key=len)
        else:
            # Starý jednoduchý formát
            mapping = map_data

        print(f"Nacten mapping z: {map_path}")
        print(f"Celkem mapovani: {len(mapping)}")
    except Exception as e:
        print(f"ERROR: Nepodarilo se nacist mapu: {e}")
        return False
    
    # Načti anonymizovaný dokument
    try:
        doc = Document(anon_doc_path)
        print(f"Nacten anonymni dokument: {anon_doc_path}")
    except Exception as e:
        print(f"ERROR: Nepodarilo se nacist dokument: {e}")
        return False
    
    print("Aplikuji deanonymizaci na dokument...")
    
    def deanonymize_text(text):
        """Nahradí tagy původními hodnotami"""
        replacements = 0
        
        # Seřaď podle délky (delší tagy první, aby se nevyměnily částečně)
        sorted_mapping = sorted(mapping.items(), key=lambda x: len(x[0]), reverse=True)
        
        for tag, original_value in sorted_mapping:
            if tag in text:
                text = text.replace(tag, original_value)
                replacements += 1
                print(f"    {tag} -> {original_value}")
        
        return text, replacements > 0
    
    total_replacements = 0
    
    # Zpracuj hlavní text
    for para in doc.paragraphs:
        for run in para.runs:
            if run.text.strip():
                new_text, changed = deanonymize_text(run.text)
                if changed:
                    run.text = new_text
                    total_replacements += 1
    
    # Zpracuj tabulky
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            new_text, changed = deanonymize_text(run.text)
                            if changed:
                                run.text = new_text
                                total_replacements += 1
    
    # Ulož deanonymizovaný dokument
    try:
        doc.save(output_path)
        print(f"Deanonymizovany dokument ulozen: {output_path}")
    except Exception as e:
        print(f"ERROR: Chyba pri ukladani: {e}")
        return False
    
    print("DEANONYMIZACE DOKONCENA!")
    print(f"Celkem nahrazeni: {total_replacements}")
    return True

# Main execution
if __name__ == "__main__":
    print(f"Vstupni soubory:")
    print(f"  Anonymni dokument: {IN_DOC}")
    print(f"  Mapa: {IN_MAP}")
    print(f"  Vystupni dokument: {OUT_DOC}")
    
    if not os.path.exists(IN_DOC):
        print(f"ERROR: Anonymni dokument neexistuje: {IN_DOC}")
        print(f"Aktualni slozka: {os.getcwd()}")
        print(f"Soubory v slozce:")
        try:
            for file in os.listdir('.'):
                if file.endswith('.docx'):
                    print(f"  - {file}")
        except:
            pass
        sys.exit(1)
    
    if not os.path.exists(IN_MAP):
        print(f"ERROR: Mapa neexistuje: {IN_MAP}")
        print(f"Aktualni slozka: {os.getcwd()}")
        print(f"JSON soubory v slozce:")
        try:
            for file in os.listdir('.'):
                if file.endswith('.json'):
                    print(f"  - {file}")
        except:
            pass
        sys.exit(1)
    
    success = deanonymize_document(IN_DOC, IN_MAP, OUT_DOC)
    
    if success:
        print("HOTOVO! Puvodni dokument byl obnoven.")
        sys.exit(0)
    else:
        print("CHYBA: Deanonymizace selhala.")
        sys.exit(1)