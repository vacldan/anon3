import sys
from pathlib import Path
import re

from docx import Document
from anon72 import Anonymizer  # for (re)anonymization if needed

sys.stdout.reconfigure(encoding="utf-8")


SENSITIVE_TYPES = {
    "PERSON",
    "ADDRESS",
    "BANK",
    "BIRTH_ID",
    "EMAIL",
    "PHONE",
    "ID_CARD",
    "ICO",
    "DIC",
}

SKIP_VALUES = {
    "praha", "brno", "ostrava", "plzeň", "liberec", "olomouc",
    "česká republika", "ceska republika", "české budějovice",
    "hradec králové", "ústí nad labem", "karlovy vary", "zlín",
    "jihlava", "pardubice", "opava", "havířov", "teplice",
    "děčín", "šumperk", "třebíč", "poděbrady",
    "manager", "services", "support", "program", "career",
    "customer", "success", "account", "risk", "web", "senior",
    "junior", "lead", "head", "team", "director", "officer",
    "consultant", "specialist", "coordinator", "developer",
    "engineer", "analyst", "architect", "compliance",
    "innovation", "innovation labs", "labs", "food", "papin food",
    "invest", "group", "company",
    "corp", "ltd", "gmbh", "inc", "digital", "solutions",
    "technology", "technologies", "software", "consulting",
    "partners", "design", "media", "studio", "agency",
    "hcp_admin", "admin",
    "nová", "nový", "nové", "nového", "novém", "novému",
}


def load_doc_text(path: Path) -> str:
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def parse_map(src: Path):
    """
    Vrátí:
      - dict[label] = {"type": type, "values": set([...])}
      - set všech původních hodnot (pro audit leaků)
    """
    base = src.with_suffix("")
    txt = base.parent / (base.name + "_map.txt")
    js = base.parent / (base.name + "_map.json")

    labels = {}
    originals = set()

    if js.exists():
        import json

        with js.open(encoding="utf-8", errors="replace") as fh:
            data = json.load(fh)
        entities = data.get("entities")
        # nové JSON mapy
        if isinstance(entities, list):
            for e in entities:
                typ = e.get("type")
                label = e.get("label")
                orig = e.get("original", "") or ""
                if not label or not typ:
                    continue
                rec = labels.setdefault(label, {"type": typ, "values": set()})
                if orig:
                    rec["values"].add(orig)
                    originals.add(orig)
        else:
            # starší formát: {TYPE: {label: {...}}}
            for typ, entries in data.items():
                if not isinstance(entries, dict):
                    continue
                for label, info in entries.items():
                    if not label.startswith("[["):
                        continue
                    variants = info.get("variants") or []
                    canonical = info.get("canonical", "")
                    rec = labels.setdefault(label, {"type": typ, "values": set()})
                    for v in variants:
                        rec["values"].add(v)
                        originals.add(v)
                    if canonical:
                        rec["values"].add(canonical)
                        originals.add(canonical)

    if txt.exists():
        with txt.open(encoding="utf-8", errors="replace") as f:
            for line in f:
                line = line.strip()
                if not line or not line.startswith("[["):
                    continue
                if "]]:" not in line:
                    continue
                label, val = line.split("]]:", 1)
                label = label + "]]"
                val = val.strip()
                # Řádky se seznamem variant začínající "- " nás nezajímají,
                # tam není typ, jen jednotlivé tvary.
                if label not in labels:
                    # typ odhadneme z prefixu
                    m = re.match(r"\[\[([A-Z_]+)_\d+\]\]", label)
                    typ = m.group(1) if m else "UNKNOWN"
                    labels[label] = {"type": typ, "values": set()}
                if val:
                    labels[label]["values"].add(val)
                    originals.add(val)

    return labels, originals


FORCE_REANON = "--force" in sys.argv


def score_contract(src: Path):
    base = src.with_suffix("")
    anon = base.parent / (base.name + "_anon.docx")

    if FORCE_REANON or not anon.exists():
        print(f"[INFO] {'Přeanoymizace' if FORCE_REANON else 'Anonymizace chybí'}: {src.name}")
        anonizator = Anonymizer()
        anonizator.anonymize_docx(str(src), str(anon), str(base) + "_map.json", str(base) + "_map.txt")

    if not anon.exists():
        return {
            "name": src.name,
            "score": 0.0,
            "max_score": 10.0,
            "errors": ["Chybí anonymizovaný soubor"],
        }

    labels, originals = parse_map(src)
    anon_text = load_doc_text(anon)

    leaks = []
    phantom = []
    tag_type_errors = []

    for v in originals:
        v_clean = v.strip()
        if len(v_clean) < 5:
            continue
        if v_clean.lower() in SKIP_VALUES:
            continue
        words = v_clean.split()
        if len(words) == 1 and len(v_clean) <= 6 and v_clean.lower() in SKIP_VALUES:
            continue
        if v_clean.isdigit() and len(v_clean) <= 10:
            if re.search(r'(?<![a-zA-Z0-9_/-])' + re.escape(v_clean) + r'(?![a-zA-Z0-9_/-])', anon_text):
                leaks.append(v_clean)
        elif v_clean in anon_text:
            leaks.append(v_clean)

    # 2) Phantom tagy – tag v anon textu, který není v mapě
    TAG_RE = re.compile(r"\[\[([A-Z_]+)_\d+\]\]")
    tags_in_anon = set(TAG_RE.findall(anon_text))
    full_tags_in_anon = set(re.findall(r"\[\[[A-Z_]+_\d+\]\]", anon_text))
    for t in full_tags_in_anon:
        if t not in labels:
            phantom.append(t)

    # 3) Základní kontrola typů – jestli typ aspoň sedí na formát hodnot
    EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
    PHONE_RE = re.compile(r"^[+0-9][0-9\s+/-]{5,}$")
    BIRTH_RE = re.compile(r"^\d{6}/?\d{3,4}$")

    for label, info in labels.items():
        typ = info["type"]
        for v in info["values"]:
            v_str = v.strip()
            if not v_str:
                continue
            if typ == "EMAIL" and not EMAIL_RE.match(v_str):
                tag_type_errors.append(f"{label}: EMAIL neodpovídá formátu ({v_str})")
            elif typ == "PHONE" and not PHONE_RE.match(v_str):
                tag_type_errors.append(f"{label}: PHONE neodpovídá formátu ({v_str})")
            elif typ == "BIRTH_ID" and not BIRTH_RE.match(v_str):
                tag_type_errors.append(f"{label}: BIRTH_ID neodpovídá formátu ({v_str})")

    # Skórování – jednoduché, konzervativní
    max_score = 10.0
    score = max_score

    if leaks:
        # tvrdý postih – únik = zásadní chyba
        score -= min(5.0, 1.0 + len(leaks) * 0.5)
    if phantom:
        score -= min(1.5, 0.3 + len(phantom) * 0.1)
    if tag_type_errors:
        score -= min(3.0, 0.5 + len(tag_type_errors) * 0.2)

    if score < 0:
        score = 0.0

    issues = []
    if leaks:
        issues.append(f"Úniky ({len(leaks)}): " + ", ".join(leaks[:5]))
    if phantom:
        issues.append(f"Phantom tagy ({len(phantom)}): " + ", ".join(phantom[:8]))
    if tag_type_errors:
        issues.append(f"Chyby typů tagů ({len(tag_type_errors)})")

    return {
        "name": src.name,
        "score": round(score, 2),
        "max_score": max_score,
        "leaks": leaks,
        "phantom": phantom,
        "tag_type_errors": tag_type_errors,
        "issues": issues,
    }


def main():
    base = Path("test_data")
    if not base.exists():
        print("test_data/ neexistuje")
        sys.exit(1)

    sources = []
    for p in sorted(base.glob("*.docx")):
        name = p.name
        if name.startswith("~$"):
            continue
        if "_anon" in name or "_deanon" in name:
            continue
        # test_pypy není plnohodnotná smlouva
        if name.startswith("test_"):
            continue
        sources.append(p)

    print(f"Nalezeno {len(sources)} zdrojových smluv v test_data/")

    results = []
    for src in sources:
        print(f"\n=== Validuji {src.name} ===")
        r = score_contract(src)
        results.append(r)
        print(f"Skóre: {r['score']}/{r['max_score']}")
        if r["issues"]:
            for iss in r["issues"]:
                print(" - " + iss)
        else:
            print(" - Bez zjištěných problémů")

    print("\n" + "=" * 70)
    print("SOUHRN VŠECH SMLUV (test_data)")
    print("=" * 70)
    print(f"{'Smlouva':<30} {'Skóre':<12} {'Úniky':<8} {'Phantom':<8} {'Typ chyby':<10}")
    print("-" * 70)
    total = 0.0
    for r in results:
        total += r["score"]
        print(
            f"{r['name']:<30} {r['score']}/{r['max_score']:<8} "
            f"{len(r['leaks']):<8} {len(r['phantom']):<8} {len(r['tag_type_errors']):<10}"
        )
    if results:
        avg = total / len(results)
        print("-" * 70)
        print(f"Průměr: {avg:.2f}/10 z {len(results)} smluv")


if __name__ == "__main__":
    main()

