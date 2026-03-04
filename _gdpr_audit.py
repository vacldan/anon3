import sys, json, re
sys.stdout.reconfigure(encoding='utf-8')
from pathlib import Path
from docx import Document

SKIP_VALUES = {'zpracovatel', 'praha', 'brno', 'ostrava', 'plzeň', 'česká republika',
               'telefon', 'e-mail', 'email', 'adresa'}

def get_text(docx_path):
    doc = Document(docx_path)
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return '\n'.join(parts)

def audit(src_path, anon_path, map_path):
    anon_text = get_text(anon_path)
    with open(map_path, encoding='utf-8') as f:
        map_data = json.load(f)

    all_originals = set()
    for e in map_data.get('entities', []):
        orig = e.get('original', '').strip()
        if orig:
            all_originals.add(orig)

    raw_leaks = []
    for v in all_originals:
        if len(v) < 4:
            continue
        if v.lower() in SKIP_VALUES:
            continue
        if v in anon_text:
            raw_leaks.append(v)

    return raw_leaks

base = Path('test_data')
results = []
for n in range(1, 51):
    src = base / f'smlouva_gdpr_test_{n:02d}.docx'
    anon = base / f'smlouva_gdpr_test_{n:02d}_anon.docx'
    mp = base / f'smlouva_gdpr_test_{n:02d}_map.json'
    if not all(p.exists() for p in [src, anon, mp]):
        continue
    leaks = audit(str(src), str(anon), str(mp))
    results.append((n, leaks))
    if leaks:
        print(f"gdpr_test_{n:02d}: {len(leaks)} LEAKS!")
        for l in leaks:
            print(f"  - {l}")

print('\n' + '=' * 60)
print('GDPR AUDIT SOUHRN (gdpr_test_01 - gdpr_test_50)')
print('=' * 60)
leaked = [n for n, l in results if l]
clean = [n for n, l in results if not l]
print(f"Clean: {len(clean)}/{len(results)}")
if leaked:
    print(f"LEAKED: {leaked}")
else:
    print("ŽÁDNÉ ÚNIKY - VŠECHNY SMLOUVY ČISTÉ!")
