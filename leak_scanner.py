"""Comprehensive leak scanner for anonymized DOCX files.
Scans all _anon.docx files for potential PII leaks."""

import os
import re
import sys
from docx import Document

test_dir = 'test_data'
anon_files = sorted([f for f in os.listdir(test_dir) if f.endswith('_anon.docx') and not f.startswith('~')])
print(f'Scanning {len(anon_files)} anonymized files for potential leaks...\n')

PSC_RE = re.compile(r'\b(\d{3}\s?\d{2})\b')
STREET_NUM_RE = re.compile(r'[A-ZÁČĎÉĚÍŇÓŘŠŤÚŮÝŽ][a-záčďéěíňóřšťúůýž]{2,}\s+\d{1,4}(?:/\d{1,4})?')

ADDRESS_CONTEXT_WORDS = [
    'sídlo', 'sídlem', 'bydliště', 'bydlištěm', 'bytem', 'adresa', 'adresy', 'adresu',
    'trvalý pobyt', 'trvalého pobytu', 'přechodný pobyt', 'místo podnikání',
    'místa podnikání', 'doručovací', 'kontaktní adres', 'korespondenční',
]

CZECH_CITIES = [
    'Praha', 'Brno', 'Ostrava', 'Plzeň', 'Liberec', 'Olomouc', 'České Budějovice',
    'Hradec Králové', 'Ústí nad Labem', 'Pardubice', 'Zlín', 'Havířov', 'Kladno',
    'Most', 'Opava', 'Frýdek-Místek', 'Karviná', 'Jihlava', 'Teplice', 'Děčín',
    'Karlovy Vary', 'Chomutov', 'Jablonec', 'Přerov', 'Prostějov', 'Třebíč',
    'Třinec', 'Tábor', 'Znojmo', 'Příbram', 'Cheb', 'Orlová', 'Kroměříž',
    'Vsetín', 'Šumperk', 'Valašské Meziříčí', 'Kolín', 'Litoměřice',
    'Mohelnice', 'Svitavy', 'Chrudim', 'Trutnov', 'Louny', 'Uherské Hradiště',
    'Břeclav', 'Hodonín', 'Vyškov', 'Blansko', 'Nový Jičín',
]

files_with_leaks = 0
total_leaks = 0

for fname in anon_files:
    path = os.path.join(test_dir, fname)
    try:
        doc = Document(path)
        full_text = '\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += '\n' + '\n'.join(p.text for p in cell.paragraphs)

        leaks = []

        # 1. Check for city names near address context keywords (not inside [[ ]])
        for city in CZECH_CITIES:
            for ctx_word in ADDRESS_CONTEXT_WORDS:
                pat = re.compile(
                    re.escape(ctx_word) + r'\s*:?\s*[^[\n]{0,30}?' + re.escape(city),
                    re.IGNORECASE
                )
                for m in pat.finditer(full_text):
                    # Check if city is NOT inside a tag
                    after_city = full_text[m.end():m.end()+5]
                    before_match = full_text[max(0,m.start()-5):m.start()]
                    if '[[' not in before_match and ']]' not in after_city:
                        context = full_text[m.start():min(len(full_text), m.end()+30)]
                        leaks.append(f'CITY near address context: "{context[:80]}"')

        # 2. Check for PSC near address context (not inside tags)
        for m in PSC_RE.finditer(full_text):
            psc_val = m.group(1).replace(' ', '')
            # Skip common non-PSC numbers
            if psc_val in ('00000', '12345', '99999', '56789'):
                continue
            before = full_text[max(0,m.start()-80):m.start()]
            after = full_text[m.end():min(len(full_text), m.end()+20)]
            # Only flag if near address context
            if any(w in before.lower() for w in ADDRESS_CONTEXT_WORDS):
                if '[[ADDRESS' not in before[-30:] and '[[ADDRESS' not in after[:20]:
                    context = full_text[max(0,m.start()-20):min(len(full_text), m.end()+30)]
                    leaks.append(f'PSC near address: "{context[:80]}"')

        # 3. Check for street+number patterns near address context
        for m in STREET_NUM_RE.finditer(full_text):
            street_val = m.group()
            before = full_text[max(0,m.start()-80):m.start()]
            if any(w in before.lower() for w in ADDRESS_CONTEXT_WORDS):
                if '[[ADDRESS' not in before[-30:] and '[[' not in full_text[m.start()-2:m.start()]:
                    # Check if street word is a common non-address word
                    first_word = street_val.split()[0].lower()
                    if first_word not in {'výše', 'částka', 'celkem', 'splatnost', 'splátka',
                                          'poplatek', 'úrok', 'dne', 'roku', 'článek', 'číslo',
                                          'verze', 'strana', 'oddíl', 'příloha', 'bod',
                                          'smlouva', 'zákon', 'jednací', 'spisová', 'podíl',
                                          'variabilní', 'specifický', 'referenční', 'osobní',
                                          'bankovní', 'objem', 'pracovní', 'zkušební',
                                          'denně', 'měsíčně', 'ročně', 'maximálně'}:
                        context = full_text[max(0,m.start()-20):min(len(full_text), m.end()+20)]
                        leaks.append(f'Street+num near address: "{context[:80]}"')

        # Deduplicate leaks
        unique_leaks = list(dict.fromkeys(leaks))

        if unique_leaks:
            files_with_leaks += 1
            total_leaks += len(unique_leaks)
            print(f'LEAKS in [{fname}] ({len(unique_leaks)} issues):')
            for l in unique_leaks[:8]:
                print(f'  {l}')
            if len(unique_leaks) > 8:
                print(f'  ... and {len(unique_leaks)-8} more')
            print()
    except Exception as e:
        print(f'ERROR [{fname}]: {e}')

print(f'\n{"="*60}')
print(f'SUMMARY: {files_with_leaks} files with potential leaks, {total_leaks} total issues')
print(f'Scanned: {len(anon_files)} files')
